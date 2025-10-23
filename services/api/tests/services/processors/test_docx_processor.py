import pytest
from pathlib import Path
from docx import Document
from docx.shared import Pt
from app.services.processors.docx_processor import DOCXProcessor


def test_docx_processor_supports_docx():
    """Test that DOCX processor supports DOCX files"""
    processor = DOCXProcessor()
    assert processor.supports_file_type("docx") is True
    assert processor.supports_file_type("pdf") is False


def test_process_simple_docx(tmp_path):
    """Test processing simple DOCX file"""
    docx_file = tmp_path / "test.docx"

    # Create a simple DOCX file
    doc = Document()
    doc.add_paragraph("Paragraph one.")
    doc.add_paragraph("Paragraph two.")
    doc.save(str(docx_file))

    processor = DOCXProcessor()
    result = processor.process(docx_file)

    assert result.document_type == "docx"
    assert len(result.elements) >= 2
    assert all(e.element_type == 'paragraph' for e in result.elements)


def test_process_docx_with_headings(tmp_path):
    """Test DOCX heading detection and hierarchy"""
    docx_file = tmp_path / "test_headings.docx"

    # Create DOCX with headings
    doc = Document()
    doc.add_heading("Main Title", level=1)
    doc.add_paragraph("Introduction paragraph.")
    doc.add_heading("Section 1", level=2)
    doc.add_paragraph("Content under section 1.")
    doc.add_heading("Subsection 1.1", level=3)
    doc.add_paragraph("Detailed content here.")
    doc.add_heading("Section 2", level=2)
    doc.add_paragraph("More content in section 2.")
    doc.save(str(docx_file))

    processor = DOCXProcessor()
    result = processor.process(docx_file)

    assert result.document_type == "docx"

    # Check that headings are detected
    headings = [e for e in result.elements if e.element_type == 'heading']
    assert len(headings) == 4  # Main Title, Section 1, Subsection 1.1, Section 2

    # Check heading levels
    assert headings[0].section_level == 1
    assert headings[1].section_level == 2
    assert headings[2].section_level == 3
    assert headings[3].section_level == 2

    # Check section hierarchy
    assert headings[0].section_title == "Main Title"
    assert headings[1].section_title == "Main Title > Section 1"
    assert headings[2].section_title == "Main Title > Section 1 > Subsection 1.1"
    assert headings[3].section_title == "Main Title > Section 2"

    # Check paragraphs have section context
    paragraphs = [e for e in result.elements if e.element_type == 'paragraph']
    assert len(paragraphs) > 0
    assert any(p.section_title is not None for p in paragraphs)


def test_process_docx_with_tables(tmp_path):
    """Test DOCX table extraction and markdown conversion"""
    docx_file = tmp_path / "test_tables.docx"

    # Create DOCX with a table
    doc = Document()
    doc.add_heading("Data Table", level=1)

    # Add a 3x3 table
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Light Grid Accent 1'

    # Header row
    header_cells = table.rows[0].cells
    header_cells[0].text = "Name"
    header_cells[1].text = "Age"
    header_cells[2].text = "City"

    # Data rows
    row1_cells = table.rows[1].cells
    row1_cells[0].text = "Alice"
    row1_cells[1].text = "30"
    row1_cells[2].text = "NYC"

    row2_cells = table.rows[2].cells
    row2_cells[0].text = "Bob"
    row2_cells[1].text = "25"
    row2_cells[2].text = "LA"

    doc.save(str(docx_file))

    processor = DOCXProcessor()
    result = processor.process(docx_file)

    # Check that table is extracted
    tables = [e for e in result.elements if e.element_type == 'table']
    assert len(tables) == 1

    # Check table content is converted to markdown
    table_content = tables[0].content
    assert "Name" in table_content
    assert "Alice" in table_content
    assert "|" in table_content  # Markdown table format
    assert "---" in table_content  # Markdown header separator

    # Check metadata
    assert result.metadata['num_tables'] == 1


def test_process_docx_metadata(tmp_path):
    """Test DOCX metadata extraction"""
    docx_file = tmp_path / "test_metadata.docx"

    doc = Document()
    doc.add_paragraph("Test paragraph 1")
    doc.add_paragraph("Test paragraph 2")
    doc.add_paragraph("Test paragraph 3")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"

    doc.save(str(docx_file))

    processor = DOCXProcessor()
    result = processor.process(docx_file)

    # Check metadata
    assert result.metadata['num_paragraphs'] == 3
    assert result.metadata['num_tables'] == 1


def test_process_docx_error_handling(tmp_path):
    """Test error handling for invalid DOCX files"""
    invalid_file = tmp_path / "invalid.docx"
    invalid_file.write_text("This is not a valid DOCX file")

    processor = DOCXProcessor()

    with pytest.raises(ValueError, match="DOCX file not found or invalid|Error processing DOCX"):
        processor.process(invalid_file)
