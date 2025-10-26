"""DOCX document processor."""
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.services.processors.base_processor import (BaseDocumentProcessor,
                                                    DocumentElement,
                                                    ProcessedDocument)


class DOCXProcessor(BaseDocumentProcessor):
    """Processor for DOCX documents."""

    def supports_file_type(self, file_type: str) -> bool:
        """Check if this processor supports the file type."""
        return file_type.lower() == "docx"

    def process(self, file_path: Path) -> ProcessedDocument:
        """
        Process DOCX and extract structured content.

        Args:
            file_path: Path to DOCX file

        Returns:
            ProcessedDocument with elements preserving structure

        Raises:
            ValueError: If DOCX is invalid, corrupt, or cannot be read
        """
        try:
            doc = Document(str(file_path))
            elements: list[DocumentElement] = []

            current_section = []
            section_level = 0

            for para in doc.paragraphs:
                if not para.text.strip():
                    continue

                # Detect headings by style
                is_heading = para.style.name.startswith("Heading")

                if is_heading:
                    # Extract heading level
                    try:
                        level = int(para.style.name.split()[-1])
                    except:
                        level = 1

                    current_section = current_section[: level - 1] + [para.text.strip()]
                    section_level = level

                    element = DocumentElement(
                        content=para.text.strip(),
                        element_type="heading",
                        metadata={"heading_level": level},
                        section_level=level,
                        section_title=" > ".join(current_section),
                    )
                else:
                    # Regular paragraph
                    element = DocumentElement(
                        content=para.text.strip(),
                        element_type="paragraph",
                        metadata={},
                        section_title=" > ".join(current_section)
                        if current_section
                        else None,
                        section_level=section_level,
                    )

                elements.append(element)

            # Extract tables
            for table_idx, table in enumerate(doc.tables):
                table_text = self._table_to_markdown(table)
                element = DocumentElement(
                    content=table_text,
                    element_type="table",
                    metadata={"table_index": table_idx},
                    section_title=" > ".join(current_section)
                    if current_section
                    else None,
                )
                elements.append(element)

            metadata = {
                "num_paragraphs": len(doc.paragraphs),
                "num_tables": len(doc.tables),
            }

        except PackageNotFoundError:
            raise ValueError(f"DOCX file not found or invalid: {file_path}")
        except FileNotFoundError:
            raise ValueError(f"DOCX file not found: {file_path}")
        except Exception as e:
            raise ValueError(f"Error processing DOCX: {e}")

        return ProcessedDocument(
            elements=elements, metadata=metadata, document_type="docx"
        )

    def _table_to_markdown(self, table) -> str:
        """Convert table to markdown format."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")

        if len(rows) > 1:
            # Add header separator
            num_cols = len(table.rows[0].cells)
            separator = "|" + "|".join(["---"] * num_cols) + "|"
            rows.insert(1, separator)

        return "\n".join(rows)
